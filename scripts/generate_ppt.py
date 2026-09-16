from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN
import os

def create_presentation():
    prs = Presentation()
    
    # 1. Title Slide
    title_slide_layout = prs.slide_layouts[0]
    slide = prs.slides.add_slide(title_slide_layout)
    title = slide.shapes.title
    subtitle = slide.placeholders[1]
    title.text = "OpsIntel"
    subtitle.text = "Executive Governance Platform\nDemo & Architecture Overview"

    # 2. What is OpsIntel?
    bullet_slide_layout = prs.slide_layouts[1]
    slide = prs.slides.add_slide(bullet_slide_layout)
    title = slide.shapes.title
    title.text = "What is OpsIntel?"
    body_shape = slide.shapes.placeholders[1]
    tf = body_shape.text_frame
    tf.text = "OpsIntel is a modern, real-time IT Service Management (ITSM) dashboard."
    
    p = tf.add_paragraph()
    p.text = "Executive 30,000-foot view of your entire IT environment."
    p.level = 1
    
    p = tf.add_paragraph()
    p.text = "Real-time analytics across Incidents, Problems, Changes, and SLAs."
    p.level = 1
    
    p = tf.add_paragraph()
    p.text = "Fast, responsive, and secure by design."
    p.level = 1

    # 3. Architecture Overview
    slide = prs.slides.add_slide(bullet_slide_layout)
    title = slide.shapes.title
    title.text = "Architecture Overview"
    body_shape = slide.shapes.placeholders[1]
    tf = body_shape.text_frame
    
    tf.text = "A Split-Stack Architecture built for scale and speed:"
    
    p = tf.add_paragraph()
    p.text = "Frontend: React.js, Vite, and Tailwind CSS"
    p.level = 1
    p = tf.add_paragraph()
    p.text = "Backend: FastAPI (Python) for high-performance REST and WebSockets"
    p.level = 1
    p = tf.add_paragraph()
    p.text = "Database: SQLite with SQLAlchemy ORM"
    p.level = 1

    # 4. Key Features
    slide = prs.slides.add_slide(bullet_slide_layout)
    title = slide.shapes.title
    title.text = "Key Features & Dashboard"
    body_shape = slide.shapes.placeholders[1]
    tf = body_shape.text_frame
    
    tf.text = "Comprehensive KPIs:"
    
    p = tf.add_paragraph()
    p.text = "Dynamic Charts: Priority distribution, Problem Aging, Change Success Rate"
    p.level = 1
    p = tf.add_paragraph()
    p.text = "Top Services: Instantly identify which services generate the most incidents"
    p.level = 1
    p = tf.add_paragraph()
    p.text = "Live Notifications: WebSockets stream alerts instantly to the dashboard"
    p.level = 1

    # 5. Service Health Tracking
    slide = prs.slides.add_slide(bullet_slide_layout)
    title = slide.shapes.title
    title.text = "Service Health Tracking"
    body_shape = slide.shapes.placeholders[1]
    tf = body_shape.text_frame
    
    tf.text = "Proprietary Operational Health Score (0-100)"
    
    p = tf.add_paragraph()
    p.text = "Auto-categorizes services into Healthy, Warning, or Critical states"
    p.level = 1
    p = tf.add_paragraph()
    p.text = "Calculated in real-time based on active P1/P2 incidents and SLA breaches"
    p.level = 1

    # 6. Data Ingestion Engine
    slide = prs.slides.add_slide(bullet_slide_layout)
    title = slide.shapes.title
    title.text = "Data Ingestion Engine"
    body_shape = slide.shapes.placeholders[1]
    tf = body_shape.text_frame
    
    tf.text = "Bulk-upload historical ITSM data effortlessly"
    
    p = tf.add_paragraph()
    p.text = "Auto-detects CSV types (Incidents, Problems, Changes, SLAs, Services)"
    p.level = 1
    p = tf.add_paragraph()
    p.text = "Upload History API tracks all uploads securely in the backend"
    p.level = 1

    # 7. Reporting and Export
    slide = prs.slides.add_slide(bullet_slide_layout)
    title = slide.shapes.title
    title.text = "Data Export & Reporting"
    body_shape = slide.shapes.placeholders[1]
    tf = body_shape.text_frame
    
    tf.text = "Empowering Executives with Data"
    
    p = tf.add_paragraph()
    p.text = "Live Paginated Tables: View up to 250 raw rows across all domains"
    p.level = 1
    p = tf.add_paragraph()
    p.text = "1-Click CSV Export: Instantly serialize table data and download locally"
    p.level = 1

    # Save presentation
    output_path = os.path.join(os.getcwd(), "OpsIntel_Presentation.pptx")
    prs.save(output_path)
    print(f"Presentation generated at {output_path}")

if __name__ == "__main__":
    create_presentation()
