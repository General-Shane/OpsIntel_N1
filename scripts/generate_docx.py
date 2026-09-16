from docx import Document
from docx.shared import Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
import os

def add_heading(doc, text, level=1):
    heading = doc.add_heading(text, level=level)
    run = heading.runs[0]
    run.font.name = 'Arial'
    run.font.color.rgb = RGBColor(0x1F, 0x49, 0x7D)  # Dark Blue

def add_paragraph(doc, text, bold=False):
    p = doc.add_paragraph()
    run = p.add_run(text)
    if bold:
        run.bold = True
    run.font.name = 'Calibri'
    run.font.size = Pt(11)
    return p

def create_document():
    doc = Document()
    
    # Title
    title = doc.add_heading('OpsIntel: Executive Governance Platform', 0)
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    doc.add_paragraph("Comprehensive Project Documentation\n").alignment = WD_ALIGN_PARAGRAPH.CENTER
    
    # 1. About the Project
    add_heading(doc, "1. About the Project", level=1)
    add_paragraph(doc, "OpsIntel is a modern, real-time IT Service Management (ITSM) and Executive Governance Platform. It is designed to provide IT leaders, Site Reliability Engineers (SREs), and Operations teams with a 30,000-foot view of their entire IT landscape. By aggregating data across multiple domains—Incidents, Problems, Changes, and Service Level Agreements (SLAs)—OpsIntel delivers actionable insights, automated health tracking, and robust data reporting capabilities.")
    
    # 2. What We Are Doing
    add_heading(doc, "2. What We Are Doing", level=1)
    add_paragraph(doc, "We are building a centralized dashboard that replaces fragmented spreadsheets and legacy ITSM tools. OpsIntel auto-ingests historical CSV data and displays complex metrics (like Priority Distributions, Problem Aging, Change Success Rates) in a single, visually stunning, real-time interface.")
    
    # 3. Why We Are Doing It
    add_heading(doc, "3. Why We Are Doing It", level=1)
    add_paragraph(doc, "Modern IT environments are highly complex. Without a unified dashboard, executives struggle to identify which services are inherently unstable (Problem Management) or if deployments are causing outages (Change Management). OpsIntel was created to solve the \"visibility gap\" by instantly highlighting SLA breaches and calculating a proprietary \"Operational Health Score\" for every service.")
    
    # 4. Tech Stack
    add_heading(doc, "4. Technology Stack", level=1)
    p = doc.add_paragraph(style='List Bullet')
    p.add_run("Frontend: ").bold = True
    p.add_run("React (TypeScript), Vite, Tailwind CSS v4, Recharts, Lucide Icons.")
    
    p = doc.add_paragraph(style='List Bullet')
    p.add_run("Backend: ").bold = True
    p.add_run("FastAPI (Python), Uvicorn, WebSockets.")
    
    p = doc.add_paragraph(style='List Bullet')
    p.add_run("Database & ORM: ").bold = True
    p.add_run("SQLite, SQLAlchemy, Pydantic.")
    
    # 5. System Architecture
    add_heading(doc, "5. System Architecture", level=1)
    add_paragraph(doc, "OpsIntel utilizes a decoupled, split-stack architecture:")
    p = doc.add_paragraph(style='List Bullet')
    p.add_run("Client Tier: ").bold = True
    p.add_run("A responsive React Single Page Application (SPA) running on Port 5173. It manages client-side routing and state, communicating with the backend exclusively via an Axios-based API client.")
    
    p = doc.add_paragraph(style='List Bullet')
    p.add_run("API & Processing Tier: ").bold = True
    p.add_run("A high-performance FastAPI server running on Port 8000. It handles JWT authentication, complex SQL aggregations for analytics, bulk CSV parsing (Data Ingestion Engine), and maintains a WebSocket server for real-time alerts.")
    
    p = doc.add_paragraph(style='List Bullet')
    p.add_run("Data Tier: ").bold = True
    p.add_run("A relational SQLite database configured with SQLAlchemy to map Python objects to tables (Incidents, Problems, Changes, SLAs, Services).")
    
    # 6. Detailed Implementation & Features
    add_heading(doc, "6. Detailed Implementation & Working", level=1)
    
    add_heading(doc, "Authentication System", level=2)
    add_paragraph(doc, "The application is secured via OAuth2 with JWT tokens. The frontend Login View POSTs to the /auth/login endpoint, which provisions a bearer token. This token is appended to all subsequent API requests.")
    
    add_heading(doc, "Dashboard & Analytics Engine", level=2)
    add_paragraph(doc, "The FastAPI backend executes optimized SQLAlchemy queries to calculate Mean Time To Resolve (MTTR), SLA Compliance Rates, and Open vs Resolved metrics. These are served via /api/v1/analytics/kpis and rendered dynamically in the UI using Recharts.")
    
    add_heading(doc, "Real-time Event Streaming", level=2)
    add_paragraph(doc, "A background scheduler injects synthetic 'live' incidents into the database every 15 seconds. Simultaneously, it broadcasts these events over WebSockets (ws://) directly to the Dashboard, triggering visual notifications instantly.")
    
    add_heading(doc, "Data Ingestion & Extraction", level=2)
    add_paragraph(doc, "Users can navigate to the Data Upload section and submit raw ITSM CSV files. The backend dynamically infers the CSV schema (e.g., matching 'problem_id' to the Problem domain) and bulk-inserts the records into SQLite. Conversely, all Table views feature a 1-Click Export utility that serializes the active table data into a clean CSV for local download.")
    
    add_heading(doc, "Service Health Algorithms", level=2)
    add_paragraph(doc, "The system calculates an Operational Health Score starting at 100. Penalties are mathematically deduced based on active P1/P2 incidents and historical SLA breaches, allowing the system to categorize services into Healthy (Green), Warning (Amber), or Critical (Red) tiers.")
    
    # Save the document
    output_path = os.path.join(os.getcwd(), "OpsIntel_Project_Report.docx")
    doc.save(output_path)
    print(f"Document generated successfully at {output_path}")

if __name__ == "__main__":
    create_document()
